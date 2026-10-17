"""Excel report: Summary, At-Risk, CO Attainment and Cleaned Data sheets.

Header styling, frozen header rows, auto-filters, number formats and
conditional formatting (colour scales for scores and fail rates, red/amber
fills for risk levels, red text for Fail) so the workbook is useful as-is.
"""

from __future__ import annotations

import io
import math

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.palette import SCORE_COLOR
from reports.data import ReportData

HEADER_FILL = PatternFill("solid", fgColor=SCORE_COLOR.lstrip("#").upper())
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=15)
SUBTITLE_FONT = Font(size=10, color="6B6963")
SECTION_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="D9D7D0")
HIGH_FILL = PatternFill("solid", fgColor="F8D4D4")
MEDIUM_FILL = PatternFill("solid", fgColor="FDEBC2")
FAIL_FONT = Font(color="C0392B", bold=True)
PCT = "0.0"
GREEN_TO_RED = ("F8696B", "FFEB84", "63BE7B")  # red (low) -> yellow -> green (high)

CLEAN_COLUMNS = {
    "roll_no": "Roll No",
    "student_name": "Name",
    "section": "Section",
    "semester": "Semester",
    "subject": "Subject",
    "attendance_pct": "Attendance %",
    "assignment": "Assignment",
    "mst1": "MST-1",
    "mst2": "MST-2",
    "end_term": "End Term",
    "internal_total": "Internal Total",
    "total": "Total",
    "score_pct": "Score %",
    "result": "Result",
    "result_reported": "Result (as in sheet)",
    "assignment_absent": "Absent: Assignment",
    "mst1_absent": "Absent: MST-1",
    "mst2_absent": "Absent: MST-2",
    "end_term_absent": "Absent: End Term",
}


def _clean_value(value):
    """Plain Python values for openpyxl: NaN/NA -> blank, numpy scalars -> Python."""
    if isinstance(value, (list, tuple, dict)):
        return "; ".join(map(str, value)) if isinstance(value, (list, tuple)) else str(value)
    if value is None or pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value


def _write_table(ws: Worksheet, frame: pd.DataFrame, start_row: int, formats: dict[str, str] | None = None) -> tuple:
    """Write a header + rows; returns (first_data_row, last_row, {column name: letter})."""
    formats = formats or {}
    letters = {}
    for j, column in enumerate(frame.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=str(column))
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
        letters[column] = get_column_letter(j)
    # Data cells get a number format only where needed: every style assignment costs openpyxl a hash
    # lookup, and per-cell borders made a 1,000-student workbook take ~10 s (see DEBUG_LOG #30).
    formatted = {j: formats[c] for j, c in enumerate(frame.columns, start=1) if c in formats}
    for i, row in enumerate(frame.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=_clean_value(value))
            if j in formatted:
                cell.number_format = formatted[j]
    return start_row + 1, start_row + len(frame), letters


def _autosize(ws: Worksheet, minimum: int = 8, maximum: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                length = len(str(cell.value)) if not isinstance(cell.value, float) else 7
                widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = max(minimum, min(maximum, width + 2))


def _summary_sheet(wb: Workbook, data: ReportData, summary: str) -> None:
    ws = wb.active
    ws.title = "Summary"
    meta = data.meta
    ws["A1"] = meta.title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = " · ".join(x for x in (meta.college, meta.department) if x) or "GradeLens report"
    ws["A2"].font = SUBTITLE_FONT
    details = [
        ("Semester", meta.semester),
        ("Faculty", meta.faculty),
        ("Scope", meta.scope),
        ("Stage", "Final results" if data.stage == "final" else "Mid-semester"),
        ("Generated", meta.generated),
    ]
    row = 4
    for label, value in details:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value or "–")
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Executive summary").font = SECTION_FONT
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=8)
    cell = ws.cell(row=row + 1, column=1, value=summary)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row + 1].height = max(60, 15 * (len(summary) // 110 + 1))
    row += 3

    k = data.kpis
    ws.cell(row=row, column=1, value="Key figures").font = SECTION_FONT
    kpis = pd.DataFrame(
        [
            ["Students", k["students"]],
            ["Records (student × subject)", k["records"]],
            ["Pass rate %", round(k["pass_rate"], 1) if not math.isnan(k["pass_rate"]) else None],
            ["Average score %", round(k["avg_score"], 1) if not math.isnan(k["avg_score"]) else None],
            ["Failed records", k["fails"]],
            ["At-risk students (high + medium)", k.get("at_risk")],
            ["Absentees", k["absentees"]],
        ],
        columns=["Measure", "Value"],
    )
    _, last, _ = _write_table(ws, kpis, row + 1)
    row = last + 2

    ws.cell(row=row, column=1, value="By section").font = SECTION_FONT
    sections = data.by_section.rename(
        columns={
            "section": "Section",
            "students": "Students",
            "records": "Records",
            "avg_score": "Avg score %",
            "pass_rate": "Pass rate %",
            "fail_rate": "Fail rate %",
            "fails": "Fails",
            "avg_attendance": "Attendance %",
        }
    )
    first, last, letters = _write_table(ws, sections, row + 1, {c: PCT for c in sections.columns if "%" in c})
    if last >= first:
        ws.conditional_formatting.add(
            f"{letters['Fail rate %']}{first}:{letters['Fail rate %']}{last}",
            ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F39A74"),
        )
    row = last + 2

    ws.cell(row=row, column=1, value="By subject (hardest first)").font = SECTION_FONT
    subjects = data.by_subject[
        ["rank", "subject", "difficulty_index", "avg_score", "fail_rate", "fails", "students", "avg_attendance"]
    ].rename(
        columns={
            "rank": "Rank",
            "subject": "Subject",
            "difficulty_index": "Difficulty",
            "avg_score": "Avg score %",
            "fail_rate": "Fail rate %",
            "fails": "Fails",
            "students": "Students",
            "avg_attendance": "Attendance %",
        }
    )
    first, last, letters = _write_table(
        ws, subjects, row + 1, {"Difficulty": PCT, **{c: PCT for c in subjects.columns if "%" in c}}
    )
    if last >= first:
        ws.conditional_formatting.add(
            f"{letters['Difficulty']}{first}:{letters['Difficulty']}{last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="EB6834"),
        )
    row = last + 2

    if not data.matrix.empty:
        ws.cell(row=row, column=1, value="Fail rate %: section × subject").font = SECTION_FONT
        matrix = data.matrix.round(1).reset_index().rename(columns={"section": "Section"})
        first, last, letters = _write_table(ws, matrix, row + 1, {c: PCT for c in matrix.columns if c != "Section"})
        start, end = get_column_letter(2), get_column_letter(len(matrix.columns))
        ws.conditional_formatting.add(
            f"{start}{first}:{end}{last}",
            ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="EB6834"),
        )
    _autosize(ws)
    ws.column_dimensions["A"].width = 34


def _risk_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("At-Risk")
    if data.risk_students is None:
        ws["A1"] = "Risk scores were not included in this report."
        return
    students = data.risk_students
    frame = pd.DataFrame(
        {
            "Roll No": students["roll_no"],
            "Name": students["student_name"] if data.meta.include_names else "(omitted)",
            "Section": students["section"],
            "Risk level": students["risk_level"],
            "Risk score": students["risk_score"].astype(int),
            "Model fail %": (students["ml_prob"] * 100).round(1),
            "Avg score %": students["avg_score"].round(1),
            "Attendance %": students["avg_attendance"].round(1),
            "Subjects failed": students["fails"],
            "Reasons": students["reasons"].map(lambda rs: "; ".join(rs) if rs else ""),
        }
    )
    if not data.meta.include_names:
        frame["Roll No"] = "(omitted)"
    if frame["Model fail %"].isna().all():
        frame = frame.drop(columns=["Model fail %"])
    ws["A1"] = "Students by risk level — decision support, not a final judgement"
    ws["A1"].font = SECTION_FONT
    first, last, letters = _write_table(ws, frame, 3, {"Avg score %": PCT, "Attendance %": PCT, "Model fail %": PCT})
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(frame.columns))}{max(last, 3)}"
    if last >= first:
        level = letters["Risk level"]
        full = f"A{first}:{get_column_letter(len(frame.columns))}{last}"
        ws.conditional_formatting.add(full, FormulaRule(formula=[f'${level}{first}="High"'], fill=HIGH_FILL))
        ws.conditional_formatting.add(full, FormulaRule(formula=[f'${level}{first}="Medium"'], fill=MEDIUM_FILL))
        score = letters["Risk score"]
        ws.conditional_formatting.add(
            f"{score}{first}:{score}{last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="D03B3B"),
        )
    _autosize(ws)
    ws.column_dimensions[letters["Reasons"]].width = 90


def _co_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("CO Attainment")
    if data.co_table is None or data.co_table.empty:
        ws["A1"] = "No CO attainment data. Open the 🎯 CO Attainment tab (question-wise marks) to include it."
        return
    s = data.co_settings
    ws["A1"] = "Course Outcome attainment"
    ws["A1"].font = SECTION_FONT
    ws["A2"] = (
        f"A student attains a CO at ≥ {s.student_threshold:g}% of its marks. Level 3 ≥ {s.level3:g}% of "
        f"students, Level 2 ≥ {s.level2:g}%, Level 1 ≥ {s.level1:g}%. Target: Level {s.target_level}."
    )
    ws["A2"].font = SUBTITLE_FONT
    co = data.co_table.assign(status=data.co_table["meets_target"].map({True: "Meets target", False: "Below target"}))
    frame = co[
        ["co", "questions", "max_marks", "students_attained", "students_assessed", "pct_attained", "level", "status"]
    ].rename(
        columns={
            "co": "CO",
            "questions": "Questions",
            "max_marks": "Max marks",
            "students_attained": "Attained",
            "students_assessed": "Assessed",
            "pct_attained": "% attained",
            "level": "Level",
            "status": "Status",
        }
    )
    first, last, letters = _write_table(ws, frame, 4, {"% attained": PCT})
    if last >= first:
        col = letters["% attained"]
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color=GREEN_TO_RED[0],
                mid_type="num",
                mid_value=s.threshold_for_level(s.target_level),
                mid_color=GREEN_TO_RED[1],
                end_type="num",
                end_value=100,
                end_color=GREEN_TO_RED[2],
            ),
        )
        status = letters["Status"]
        ws.conditional_formatting.add(
            f"{status}{first}:{status}{last}", CellIsRule(operator="equal", formula=['"Below target"'], font=FAIL_FONT)
        )
    _autosize(ws)


def _data_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Cleaned Data")
    columns = [c for c in CLEAN_COLUMNS if c in data.clean.columns]
    frame = data.clean[columns].rename(columns=CLEAN_COLUMNS)
    if not data.meta.include_names:
        frame = frame.drop(columns=[c for c in ("Name",) if c in frame])
        frame["Roll No"] = "(omitted)"
    # Thousands of rows: round instead of styling each cell, and append rows in bulk.
    floats = [c for c in frame.columns if pd.api.types.is_float_dtype(frame[c])]
    frame = frame.assign(**{c: frame[c].round(2) for c in floats})
    _, _, letters = _write_table(ws, frame.head(0), 1)
    for row in frame.itertuples(index=False):
        ws.append([_clean_value(v) for v in row])
    first, last = 2, 1 + len(frame)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{max(last, 1)}"
    if last >= first:
        if "Score %" in letters:
            col = letters["Score %"]
            ws.conditional_formatting.add(
                f"{col}{first}:{col}{last}",
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color=GREEN_TO_RED[0],
                    mid_type="num",
                    mid_value=data.config.pass_pct + 10,
                    mid_color=GREEN_TO_RED[1],
                    end_type="num",
                    end_value=100,
                    end_color=GREEN_TO_RED[2],
                ),
            )
        if "Result" in letters:
            col = letters["Result"]
            ws.conditional_formatting.add(
                f"{col}{first}:{col}{last}", CellIsRule(operator="equal", formula=['"Fail"'], font=FAIL_FONT)
            )
    for column, letter in letters.items():
        ws.column_dimensions[letter].width = max(9, min(28, len(column) + 3))
    if "Name" in letters:
        ws.column_dimensions[letters["Name"]].width = 24


def _neutralise_formulas(wb: Workbook) -> None:
    """Store text that starts with "=" as text, never as a formula.

    openpyxl turns any string beginning with "=" into a live formula, so a student named
    '=HYPERLINK("http://…")' in an uploaded sheet would become a clickable link in the HOD's
    workbook (formula injection, DEBUG_LOG #39). GradeLens never writes real formulas.
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def build_excel(data: ReportData, summary: str) -> bytes:
    wb = Workbook()
    _summary_sheet(wb, data, summary)
    _risk_sheet(wb, data)
    _co_sheet(wb, data)
    _data_sheet(wb, data)
    _neutralise_formulas(wb)
    wb.properties.title = data.meta.title
    wb.properties.creator = "GradeLens"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
