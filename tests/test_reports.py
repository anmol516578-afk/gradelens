from __future__ import annotations

import io
import time

import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from agent.llm import MockLLM, text_reply
from core.co_attainment import COSettings, compute_co_attainment, parse_question_marks
from ml import risk_model as rm
from reports.data import ReportMeta, build_report_data, key_facts, llm_summary, template_summary
from reports.excel_report import build_excel
from reports.pdf_report import build_pdf

META = ReportMeta(
    college="Rayat Bahra University",
    department="Computer Applications",
    semester="5",
    faculty="Prof. Test",
    scope="All students",
)


@pytest.fixture(scope="module")
def risk(derived, config):
    return rm.assess(derived, config, rm.train(derived))


@pytest.fixture(scope="module")
def co(bundle):
    settings = COSettings()
    table, _, _ = compute_co_attainment(parse_question_marks(bundle.co_marks), bundle.co_mapping, settings)
    return {"table": table, "settings": settings}


@pytest.fixture(scope="module")
def data(derived, config, risk, co, prepared_long):
    return build_report_data(derived, config, META, risk=risk, co_result=co, quality=prepared_long[0].report)


def pdf_text(pdf: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf))
    return "\n".join(page.extract_text() for page in reader.pages)


def test_template_summary_uses_real_numbers(data):
    text = template_summary(data)
    assert "180 students across 5 subjects" in text
    assert "pass rate is 89.4%" in text
    assert "Discrete Maths is the most difficult subject" in text
    assert "Section B needs the most attention" in text
    assert "CO3 falls below the CO attainment target" in text


def test_pdf_is_valid_and_complete(data):
    pdf = build_pdf(data, "Summary paragraph for the HOD ≥ test.")
    assert pdf.startswith(b"%PDF") and len(pdf) > 50_000
    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) >= 3
    assert reader.metadata.title == "Student Performance Report"
    text = pdf_text(pdf)
    for expected in (
        "Student Performance Report",
        "Rayat Bahra University",
        "Executive summary",
        "Summary paragraph for the HOD ≥ test.",
        "Results by section",
        "Results by subject",
        "Discrete Maths",
        "Students who may need support",
        "Course Outcome attainment",
        "Method and data notes",
        "Prof. Test",
    ):
        assert expected in text, expected
    top_risk = data.at_risk.iloc[0]
    assert top_risk["student_name"] in text


def test_pdf_without_names_hides_identities(derived, config, risk, co):
    meta = ReportMeta(scope="All students", include_names=False)
    data = build_report_data(derived, config, meta, risk=risk, co_result=co)
    text = pdf_text(build_pdf(data, "Summary."))
    assert "Student names are omitted" in text
    assert not any(name in text for name in derived["student_name"].unique())
    assert "2023BCA" not in text


def test_pdf_without_risk_or_co(derived, config):
    data = build_report_data(derived, config, ReportMeta(scope="Section B"))
    text = pdf_text(build_pdf(data, "Summary."))
    assert "Students who may need support" not in text and "Course Outcome" not in text


def test_excel_sheets_styles_and_conditional_formatting(data, derived):
    xlsx = build_excel(data, "Summary.")
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Summary", "At-Risk", "CO Attainment", "Cleaned Data"]
    summary = wb["Summary"]
    assert summary["A1"].value == "Student Performance Report"
    assert summary["A1"].font.bold

    at_risk = wb["At-Risk"]
    header = [c.value for c in at_risk[3]]
    assert header[:5] == ["Roll No", "Name", "Section", "Risk level", "Risk score"]
    assert at_risk["A3"].font.bold and at_risk["A3"].fill.fgColor.rgb.upper().endswith("2A78D6")
    assert at_risk.max_row - 3 == len(data.risk_students) == 180  # everyone, not just the PDF's top 40
    assert at_risk.freeze_panes == "A4" and at_risk.auto_filter.ref
    assert len(at_risk.conditional_formatting) >= 2

    co = wb["CO Attainment"]
    assert [c.value for c in co[4]][:2] == ["CO", "Questions"]
    assert co.max_row - 4 == len(data.co_table)

    clean = wb["Cleaned Data"]
    assert clean.max_row - 1 == len(derived)
    assert clean.freeze_panes == "A2"
    assert len(clean.conditional_formatting) >= 2
    # Absent marks stay blank, not zero.
    headers = [c.value for c in clean[1]]
    mst1 = headers.index("MST-1") + 1
    absent = headers.index("Absent: MST-1") + 1
    absent_rows = [r for r in range(2, clean.max_row + 1) if clean.cell(r, absent).value is True]
    assert absent_rows and all(clean.cell(r, mst1).value is None for r in absent_rows)


def test_excel_without_names(derived, config, risk):
    data = build_report_data(derived, config, ReportMeta(include_names=False), risk=risk)
    wb = load_workbook(io.BytesIO(build_excel(data, "Summary.")))
    blob = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert not any(name in blob for name in derived["student_name"].unique()[:50])
    assert "2023BCA001" not in blob
    assert "No CO attainment data" in wb["CO Attainment"]["A1"].value


def test_reports_build_fast_for_1000_students(derived, config):
    """Spec: generation under 10 seconds for 1,000 students."""
    copies = []
    for i in range(6):  # 180 x 6 = 1,080 students, 5,400 records
        copy = derived.copy()
        copy["roll_no"] = copy["roll_no"] + f"-{i}"
        copy["section"] = copy["section"] + str(i)
        copies.append(copy)
    big = pd.concat(copies, ignore_index=True)
    risk = rm.assess(big, config)
    start = time.perf_counter()
    data = build_report_data(big, config, META, risk=risk)
    pdf = build_pdf(data, template_summary(data))
    xlsx = build_excel(data, template_summary(data))
    elapsed = time.perf_counter() - start
    assert pdf and xlsx
    assert elapsed < 10, f"reports took {elapsed:.1f}s"


def test_llm_summary_sends_only_aggregates(data):
    mock = MockLLM([text_reply("Results are solid overall; focus support on Section B and Discrete Maths.")])
    text = llm_summary(mock, data)
    assert text.startswith("Results are solid")
    payload = mock.payload_text()
    assert "89.4%" in payload and "Discrete Maths" in payload
    assert not any(name in payload for name in data.clean["student_name"].unique())
    assert "2023BCA" not in payload


def test_llm_summary_rejects_empty_reply(data):
    with pytest.raises(ValueError, match="empty"):
        llm_summary(MockLLM([text_reply("   ")]), data)


def test_key_facts_cover_main_findings(data):
    facts = " ".join(key_facts(data))
    for expected in (
        "180 students",
        "Pass rate 89.4%",
        "Hardest subject: Discrete Maths",
        "Section B",
        "At-risk students",
        "CO attainment",
    ):
        assert expected in facts, expected
