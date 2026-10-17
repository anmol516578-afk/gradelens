"""Messy real-world sheets through the whole app, plus hostile-input security checks.

Every sheet must either work end to end (cleaning, analytics, risk, reports, masking, every tab)
or be rejected with a friendly message. Hostile values must never become live HTML or formulas.
"""

from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from agent.masking import Pseudonymizer
from core import analytics as an
from core.column_mapping import MappingError
from core.config import AnalysisConfig
from core.data_loader import LoadError, read_table
from core.derive import add_derived
from core.pipeline import auto_prepare
from ml import risk_model as rm
from reports.data import ReportMeta, build_report_data, template_summary
from reports.excel_report import build_excel
from reports.pdf_report import build_pdf
from tests.conftest import ROOT
from tests.messy_sheets import HTML_SUBJECT, INJECTION_NAME, SHEETS, UNREADABLE
from ui import components as C

READABLE = [name for name in SHEETS if name not in UNREADABLE]


def run_pipeline(name: str):
    filename, data = SHEETS[name]
    raw = read_table(data, filename)
    prepared, config, _ = auto_prepare(raw, AnalysisConfig())
    derived = add_derived(prepared.clean, config)
    ok, reason = rm.can_train(derived)
    risk = rm.assess(derived, config, rm.train(derived) if ok else None, fallback_reason=reason)
    report = build_report_data(derived, config, ReportMeta(), risk=risk, quality=prepared.report)
    return derived, risk, report


@pytest.mark.parametrize("name", READABLE)
def test_messy_sheet_works_end_to_end(name):
    derived, risk, report = run_pipeline(name)
    assert len(derived) > 0 and derived["roll_no"].nunique() > 0
    assert derived["score_pct"].dropna().between(0, 100).all()
    for call in (an.kpis, an.student_summary, an.subject_difficulty, an.fail_rate_matrix, an.mst_change):
        call(derived)
    an.student_profile(derived, derived["roll_no"].iloc[0])
    assert len(risk.students) == derived["roll_no"].nunique()
    assert build_pdf(report, template_summary(report)).startswith(b"%PDF")
    assert build_excel(report, "Summary")[:2] == b"PK"
    pseudo = Pseudonymizer.from_frame(derived)
    blob = pseudo.mask_frame(derived).to_csv()
    assert not [n for n in pseudo.real_names() if len(n) > 3 and n in blob]


@pytest.mark.parametrize("name", sorted(UNREADABLE))
def test_unusable_sheet_gets_a_friendly_message(name):
    with pytest.raises((LoadError, MappingError)) as info:
        run_pipeline(name)
    assert len(str(info.value)) > 10


def test_counts_on_awkward_layouts():
    assert run_pipeline("xlsx_numeric_rolls")[0]["roll_no"].nunique() == 40  # numeric rolls, one per student
    assert run_pipeline("wide_subject_codes")[0]["subject"].nunique() == 2
    assert run_pipeline("multi_semester")[0]["semester"].nunique() == 2
    special = run_pipeline("special_characters")[0]
    assert "DBMS & SQL <Lab>" in set(special["subject"])
    assert "Zoë Ångström" in set(special["student_name"])
    latin = run_pipeline("latin1_semicolon")[0]
    assert latin["student_name"].iloc[0] == "Jürgen Müller"


def test_csv_title_rows_and_ragged_rows_are_read():
    title = b"Rayat Bahra University - MST Marks\n\nRoll No,Name,MST 1\n1,Asha,12\n2,Ravi,9\n"
    assert list(read_table(title, "t.csv").columns) == ["Roll No", "Name", "MST 1"]
    ragged = b"Roll No,Name,MST 1\n1,Asha,12,extra\n2,Ravi,9\n"
    assert read_table(ragged, "r.csv").shape == (2, 4)
    semicolons = "Roll No;Name;MST 1\n1;Rao, Asha;12\n".encode("latin-1")
    assert read_table(semicolons, "s.csv").iloc[0].tolist() == ["1", "Rao, Asha", "12"]


def test_class_wide_absence_is_a_data_issue_not_40_warnings():
    filename, data = SHEETS["all_absent_mst1"]
    prepared, _, _ = auto_prepare(read_table(data, filename), AnalysisConfig())
    assert any(issue.category == "Class-wide absence" for issue in prepared.report.issues)
    _, risk, _ = run_pipeline("all_absent_mst1")
    assert not any(any(r.startswith("Absent in") for r in reasons) for reasons in risk.students["reasons"])


# --- Hostile values -------------------------------------------------------------------


def test_excel_report_never_contains_live_formulas():
    _, _, report = run_pipeline("hostile_values")
    wb = load_workbook(io.BytesIO(build_excel(report, "=1+1")))
    cells = [c for ws in wb for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert any(c.value == INJECTION_NAME for c in cells)  # kept, as plain text


def test_csv_downloads_neutralise_formulas():
    frame = pd.DataFrame({"name": [INJECTION_NAME, "+SUM(1,2)", "@cmd", "-7", "Asha"], "score": [1, 2, 3, 4, 5]})
    safe = C.csv_safe(frame)
    assert safe["name"].tolist() == ["'" + INJECTION_NAME, "'+SUM(1,2)", "'@cmd", "-7", "Asha"]
    assert safe["score"].tolist() == [1, 2, 3, 4, 5]


def test_pdf_shows_special_characters_literally():
    sheet = pd.DataFrame(
        {
            "Roll No": [f"R{i}" for i in range(6)],
            "Name": [f"N {i}" for i in range(6)],
            "Section": ["A&B"] * 6,
            "Subject": ["C++ & DS"] * 6,
            "Total": [30 + 10 * i for i in range(6)],
        }
    )
    raw = read_table(sheet.to_csv(index=False).encode(), "s.csv")
    prepared, config, _ = auto_prepare(raw, AnalysisConfig())
    derived = add_derived(prepared.clean, config)
    report = build_report_data(derived, config, ReportMeta(), risk=rm.assess(derived, config))
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(report, "x"))).pages)
    assert "A&B" in text and "C++ & DS" in text
    assert "&amp;" not in text


def _app_with(name: str) -> AppTest:
    filename, data = SHEETS[name]
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.sidebar.file_uploader[0].upload(filename, data).run()
    return at


@pytest.mark.parametrize(
    "name",
    [
        "hostile_values",
        "special_characters",
        "only_totals",
        "one_student",
        "wide_subject_codes",
        "multi_semester",
        "everyone_fails",
        "all_absent_mst1",
    ],
)
def test_every_tab_renders_for_awkward_sheets(name, monkeypatch):
    for var in ("LLM_PROVIDER", "LLM_API_KEY", "LLM_MODEL", "LLM_BASE_URL"):
        monkeypatch.setenv(var, "")
    at = _app_with(name)
    assert not at.exception, [e.value for e in at.exception]
    for label in C.TABS:
        at.session_state[C.TAB_KEY] = label
        at.run()
        assert not at.exception, (label, [e.value for e in at.exception])
        assert not at.error, (label, [e.value for e in at.error])


def test_html_in_sheet_values_is_escaped_in_chips():
    at = _app_with("hostile_values")
    at.multiselect(key="f_subjects").set_value([HTML_SUBJECT]).run()
    chip_html = " ".join(m.value for m in at.markdown if "gl-chip" in m.value)
    assert "<img" not in chip_html and "&lt;img" in chip_html


@pytest.mark.parametrize("name", ["zero_bytes", "header_only", "not_a_table"])
def test_unusable_uploads_show_a_message_in_the_app(name):
    at = _app_with(name)
    assert not at.exception
    assert at.error, "the user must be told why nothing loaded"
